# -*- coding: utf-8 -*-
"""Native (editable) PowerPoint charts with a consistent look across the deck."""
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK, XL_TICK_LABEL_POSITION
from pptx.util import Inches, Pt
from kit import rgb, INK, SLATE, BORDER, BODY


def _base(chart, size=12, legend=False, legend_pos=XL_LEGEND_POSITION.BOTTOM):
    chart.has_title = False
    chart.font.size = Pt(size)
    chart.font.name = BODY
    chart.font.color.rgb = rgb(INK)
    chart.has_legend = legend
    if legend:
        chart.legend.position = legend_pos
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)


def _cat_axis(chart, size=12, color=INK, line=BORDER):
    ca = chart.category_axis
    ca.has_major_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(size)
    ca.tick_labels.font.color.rgb = rgb(color)
    ca.format.line.color.rgb = rgb(line)
    return ca


def _hide_val_axis(chart, vmin=None, vmax=None):
    va = chart.value_axis
    va.visible = False
    va.has_major_gridlines = False
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    return va


def column(slide, x, y, w, h, cats, values, colors, fmt='#,##0', vmax=None, size=12, label_size=14, gap=55,
           name=None, alt=None):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("Series", values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    _base(ch, size)
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.vary_by_categories = False
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(label_size)
    dl.font.bold = True
    dl.font.color.rgb = rgb(INK)
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
    _cat_axis(ch, size)
    _hide_val_axis(ch, 0, vmax)
    if name:
        gf.name = name
    if alt:
        gf._element.xpath(".//p:cNvPr")[0].set("descr", alt)
    return gf


def stacked_bar(slide, x, y, w, h, cat, parts, fmt='#,##0" kWh"', size=12, label_size=14, name=None, alt=None):
    """One horizontal stacked bar. parts: [(series_name, value, fill_hex, label_hex)]"""
    cd = CategoryChartData()
    cd.categories = [cat]
    for n, v, _, _ in parts:
        cd.add_series(n, [v])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    _base(ch, size, legend=True)
    plot = ch.plots[0]
    plot.gap_width = 25
    plot.overlap = 100
    for i, (n, v, fill, lab) in enumerate(parts):
        ser = plot.series[i]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(fill)
        ser.data_labels.show_value = True
        ser.data_labels.number_format = fmt
        ser.data_labels.number_format_is_linked = False
        ser.data_labels.position = XL_LABEL_POSITION.CENTER
        ser.data_labels.font.size = Pt(label_size)
        ser.data_labels.font.bold = True
        ser.data_labels.font.color.rgb = rgb(lab)
    ca = ch.category_axis
    ca.visible = False
    _hide_val_axis(ch, 0)
    if name:
        gf.name = name
    if alt:
        gf._element.xpath(".//p:cNvPr")[0].set("descr", alt)
    return gf


def waterfall(slide, x, y, w, h, cats, bases, vals, colors, labels, vmax, size=12, label_size=13, name=None, alt=None, label_colors=None):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("base", bases)
    cd.add_series("value", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    _base(ch, size)
    plot = ch.plots[0]
    plot.gap_width = 40
    plot.overlap = 100
    base, val = plot.series[0], plot.series[1]
    base.format.fill.background()
    base.format.line.fill.background()
    for i, c in enumerate(colors):
        pt = val.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
        dlab = pt.data_label
        tf = dlab.text_frame
        tf.text = labels[i]
        r = tf.paragraphs[0].runs[0]
        r.font.size = Pt(label_size)
        r.font.bold = True
        r.font.color.rgb = rgb((label_colors or [INK] * len(colors))[i])
        dlab.position = XL_LABEL_POSITION.INSIDE_END if False else XL_LABEL_POSITION.INSIDE_BASE
    _cat_axis(ch, size)
    _hide_val_axis(ch, 0, vmax)
    if name:
        gf.name = name
    if alt:
        gf._element.xpath(".//p:cNvPr")[0].set("descr", alt)
    return gf


def line(slide, x, y, w, h, cats, values, color, fmt, vmin, vmax, size=12, label_size=13, name=None, alt=None):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("Cumulative cash flow", values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    _base(ch, size)
    plot = ch.plots[0]
    ser = plot.series[0]
    ser.smooth = False
    ser.format.line.color.rgb = rgb(color)
    ser.format.line.width = Pt(3)
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = rgb(color)
    ser.marker.format.line.color.rgb = rgb(color)
    ser.marker.size = 9
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.ABOVE
    dl.font.size = Pt(label_size)
    dl.font.bold = True
    dl.font.color.rgb = rgb(INK)
    ca = _cat_axis(ch, size)
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale = vmin, vmax
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(BORDER)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(size)
    va.tick_labels.font.color.rgb = rgb(SLATE)
    va.tick_labels.number_format = '"₹"#,##0"k";"−₹"#,##0"k"'
    va.tick_labels.number_format_is_linked = False
    if name:
        gf.name = name
    if alt:
        gf._element.xpath(".//p:cNvPr")[0].set("descr", alt)
    return gf
