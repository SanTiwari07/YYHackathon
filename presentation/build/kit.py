# -*- coding: utf-8 -*-
"""Small python-pptx toolkit for the AgroStruxure deck: native shapes, text, tables and charts.

All coordinates are inches on a 13.333 x 7.5 in (16:9) slide. Fonts are Calibri / Consolas so
PowerPoint renders with the same metrics on a judge's machine as it does here.
"""
import copy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SLIDE_W, SLIDE_H = 13.333, 7.5
BODY, MONO, DEVA = "Calibri", "Consolas", "Nirmala UI"

# ---- palette (brand tokens from docs/DESIGN_SYSTEM.md) -------------------------------------
FOREST = "024230"      # Yuva Yodha deep forest (dominant)
FOREST_2 = "06553E"    # lifted forest for cards on dark slides
EMERALD = "0D8752"
GREEN = "3DCD58"       # Schneider green
GOLD = "FFCE00"        # solar gold (fills only on light slides)
GOLD_TXT = "8A6200"    # gold usable as text on light backgrounds
WATER = "0087CD"
COLD = "7A18D8"
ALARM = "C40A0A"
AMBER = "B35F00"
INK = "10201A"
SLATE = "4F6159"       # muted text on light (>= 6:1 on white)
MIST = "F2F6F3"       # light slide background
CARD = "FFFFFF"
BORDER = "D5E0D9"
TINT_GREEN, TINT_BLUE, TINT_GOLD, TINT_VIOLET, TINT_RED = "E4F5E9", "E2F1FA", "FFF4C9", "EFE4FB", "FCE6E6"
WHITE = "FFFFFF"
ON_DARK_MUTED = "BFD8CB"


def rgb(h):
    return RGBColor.from_string(h)


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def blank(prs, bg=MIST):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = rgb(bg)
    return s


def _name(shape, name, alt):
    if name:
        shape.name = name
    if alt:
        shape._element.xpath(".//p:cNvPr")[0].set("descr", alt)


def _apply_run(run, size, color, bold, italic, font):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = rgb(color)
    if font == DEVA:  # complex-script font so Devanagari renders in Nirmala UI
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:cs", "a:ea"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", font)


def _fill_tf(tf, paras, size, color, bold, italic, font, align, line_spacing, space_after):
    """paras: str | list of (str | list-of-runs | dict).  run: str | (str, {opts})."""
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        popts = {}
        if isinstance(p, dict):
            popts = {k: v for k, v in p.items() if k != "runs"}
            p = p["runs"]
        if isinstance(p, (str, tuple)):
            p = [p]
        al = popts.get("align", align)
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[al]
        ls = popts.get("line_spacing", line_spacing)
        if ls:
            para.line_spacing = ls
        sa = popts.get("space_after", space_after)
        if sa is not None:
            para.space_after = Pt(sa)
        for r in p:
            if isinstance(r, str):
                r = (r, {})
            txt, o = r
            run = para.add_run()
            run.text = txt
            _apply_run(run, o.get("size", popts.get("size", size)), o.get("color", popts.get("color", color)),
                       o.get("bold", popts.get("bold", bold)), o.get("italic", popts.get("italic", italic)),
                       o.get("font", popts.get("font", font)))


def text(slide, x, y, w, h, paras, size=14, color=INK, bold=False, italic=False, font=BODY,
         align="l", anchor="t", line_spacing=1.0, space_after=None, name=None, alt=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    _fill_tf(tf, paras, size, color, bold, italic, font, align, line_spacing, space_after)
    _name(tb, name, alt)
    return tb


def rect(slide, x, y, w, h, fill=CARD, line=None, line_w=0.75, radius=None, shadow=False, name=None, alt=None,
         shape=None, paras=None, size=14, color=INK, bold=False, italic=False, font=BODY, align="c", anchor="m",
         pad=0.08, line_spacing=1.0, space_after=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(line_w)
    if not shadow:  # kill theme shadow
        spPr = sp._element.spPr
        if spPr.find(qn("a:effectLst")) is None:
            etree.SubElement(spPr, qn("a:effectLst"))
    else:
        spPr = sp._element.spPr
        eff = spPr.find(qn("a:effectLst"))
        if eff is None:
            eff = etree.SubElement(spPr, qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="76200", dist="25400", dir="5400000", algn="t", rotWithShape="0")
        clr = etree.SubElement(sh, qn("a:srgbClr"), val="0B2A22")
        etree.SubElement(clr, qn("a:alpha"), val="14000")
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(pad * 0.6)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if paras is not None:
        _fill_tf(tf, paras, size, color, bold, italic, font, align, line_spacing, space_after)
    _name(sp, name, alt)
    return sp


def pill(slide, x, y, w, h, label, fill, color, size=11, bold=True, name=None):
    return rect(slide, x, y, w, h, fill=fill, radius=h / 2, paras=label, size=size, color=color, bold=bold,
                align="c", anchor="m", pad=0.06, name=name)


def circle(slide, x, y, d, label, fill, color=WHITE, size=14, bold=True, name=None):
    return rect(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL, paras=label, size=size, color=color,
                bold=bold, align="c", anchor="m", pad=0, name=name)


def arrow(slide, x1, y1, x2, y2, color=SLATE, width=2.0, head=True, dash=False, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn("a:prstDash"), val="dash")
    if head:
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    if name:
        c.name = name
    return c


def image(slide, path, x, y, w=None, h=None, alt=None, name=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
    _name(pic, name, alt)
    return pic


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


# ---- tables ---------------------------------------------------------------------------------
def table(slide, x, y, w, col_w, rows, row_h=0.42, header_fill=FOREST, header_color=WHITE, size=14,
          zebra=(CARD, "F7FAF8"), col_align=None, bold_first_col=True, name=None, cell_colors=None):
    nrows, ncols = len(rows), len(rows[0])
    gs = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = gs.table
    # strip default style banding
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid" -> we override borders anyway
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i in range(nrows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            val = rows[i][j]
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            is_h = i == 0
            col = header_color if is_h else INK
            if cell_colors and (i, j) in cell_colors:
                col = cell_colors[(i, j)]
            al = (col_align[j] if col_align else "l")
            _fill_tf(tf, [val], size, col, is_h or (bold_first_col and j == 0), False,
                     MONO if (not is_h and j > 0 and col_align and col_align[j] == "r") else BODY, al, 1.0, None)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if is_h else zebra[i % 2])
            # thin bottom border
            tcPr = cell._tc.get_or_add_tcPr()
            for side in ("a:lnL", "a:lnR", "a:lnT"):
                ln = etree.SubElement(tcPr, qn(side), w="0")
                etree.SubElement(ln, qn("a:noFill"))
            ln = etree.SubElement(tcPr, qn("a:lnB"), w="9525")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=BORDER)
            # schema order: borders must precede fill inside tcPr
            fill_el = tcPr.find(qn("a:solidFill"))
            if fill_el is not None:
                tcPr.remove(fill_el)
                tcPr.append(fill_el)
    if name:
        gs.name = name
    return gs


# ---- standard slide furniture -----------------------------------------------------------------
LOGO = None  # set by build script


def header(slide, kicker, title_runs, subtitle=None, num=None, dark=False, kicker_w=3.6):
    fg = WHITE if dark else FOREST
    sub = ON_DARK_MUTED if dark else SLATE
    pill(slide, 0.6, 0.42, kicker_w, 0.32, kicker, fill=(FOREST_2 if dark else TINT_GREEN),
         color=(GREEN if dark else EMERALD), size=11, name="Kicker")
    if LOGO:
        if dark:
            rect(slide, 11.52, 0.34, 1.21, 0.58, fill=WHITE, radius=0.08, name="Logo plate")
        image(slide, LOGO[0], 11.58, 0.38, w=1.09, alt="Schneider Electric logo", name="Schneider logo")
    text(slide, 0.6, 0.86, 12.1, 0.62, [title_runs] if isinstance(title_runs, list) else title_runs,
         size=32, bold=True, color=fg, name="Title")
    if subtitle:
        text(slide, 0.6, 1.52, 11.9, 0.62, subtitle, size=16, color=sub, name="Subtitle")


def footer(slide, source, num, total=10, dark=False):
    c = ON_DARK_MUTED if dark else SLATE
    text(slide, 0.6, 6.88, 10.5, 0.5, source, size=10, color=c, name="Sources")
    text(slide, 11.2, 6.88, 1.53, 0.3, f"AgroStruxure™  ·  {num:02d} / {total}", size=10, color=c, align="r",
         name="Slide number")


def card(slide, x, y, w, h, fill=CARD, line=BORDER, radius=0.12, shadow=True, name=None):
    return rect(slide, x, y, w, h, fill=fill, line=line, radius=radius, shadow=shadow, name=name)
