# -*- coding: utf-8 -*-
"""
AgroStruxure™ Export Pipeline:
1. Generates 10 single-slide HTML pages.
2. Captures 1920x1080 PNG screenshots using headless Chrome.
3. Compiles high-resolution AgroStruxure_YuvaYodha_2026_Final.pdf.
4. Generates AgroStruxure_YuvaYodha_2026_Final.pptx using python-pptx.
"""

import os
import subprocess
from PIL import Image
import pptx
from pptx.util import Inches

def main():
    chrome_path = r'C:\Program Files\Google\Chrome\Application\chrome.exe'
    base_dir = r'D:\Research Work\YuvaYodhaHackathon Research'
    html_path = os.path.join(base_dir, 'presentation.html')
    output_dir = os.path.join(base_dir, 'rendered_slides')
    os.makedirs(output_dir, exist_ok=True)

    with open(html_path, 'r', encoding='utf-8') as f:
        html = f.read()

    png_files = []

    print('Step 1: Rendering 10 high-resolution slide screenshots via headless Chrome...')
    for i in range(1, 11):
        # Override active slide
        slide_css = f"""
    #slide-{i} {{ display: flex !important; opacity: 1 !important; z-index: 10; }}
    .slide {{ display: none !important; }}
    #slide-{i} {{ display: flex !important; }}
"""
        slide_html = html.replace('</style>', f'{slide_css}\n  </style>')
        slide_file = os.path.join(output_dir, f'slide_{i:02d}.html')
        with open(slide_file, 'w', encoding='utf-8') as sf:
            sf.write(slide_html)

        png_file = os.path.join(output_dir, f'slide_{i:02d}.png')
        file_url = 'file:///' + slide_file.replace('\\', '/')
        
        cmd = [
            chrome_path,
            '--headless',
            '--disable-gpu',
            '--no-sandbox',
            '--hide-scrollbars',
            '--window-size=1920,1080',
            f'--screenshot={png_file}',
            file_url
        ]
        
        res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        if os.path.exists(png_file) and os.path.getsize(png_file) > 10000:
            png_files.append(png_file)
            print(f'  [OK] Rendered slide_{i:02d}.png ({os.path.getsize(png_file):,} bytes)')
        else:
            print(f'  [ERROR] Failed to render slide_{i:02d}.png: {res.stderr.decode("utf-8", errors="ignore")}')

    print(f'\nStep 2: Compiling PDF from {len(png_files)} PNG slides...')
    if png_files:
        pdf_path = os.path.join(base_dir, 'AgroStruxure_YuvaYodha_2026_Final.pdf')
        images = [Image.open(p).convert('RGB') for p in png_files]
        images[0].save(pdf_path, save_all=True, append_images=images[1:], resolution=150.0, quality=95)
        print(f'  [OK] Successfully compiled PDF: {pdf_path} ({os.path.getsize(pdf_path):,} bytes)')

    print('\nStep 3: Generating native PowerPoint presentation (AgroStruxure_YuvaYodha_2026_Final.pptx)...')
    prs = pptx.Presentation()
    # Set 16:9 widescreen dimensions (13.333 x 7.5 inches = 1920x1080 aspect ratio)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    for idx, png_file in enumerate(png_files, start=1):
        slide = prs.slides.add_slide(blank_layout)
        # Add high-resolution rendered graphic covering full 16:9 slide
        slide.shapes.add_picture(png_file, 0, 0, width=prs.slide_width, height=prs.slide_height)
        print(f'  [OK] Added Slide {idx:02d} to PPTX')

    pptx_path = os.path.join(base_dir, 'AgroStruxure_YuvaYodha_2026_Final.pptx')
    prs.save(pptx_path)
    print(f'  [OK] Successfully saved PowerPoint: {pptx_path} ({os.path.getsize(pptx_path):,} bytes)')

    print('\nExport Pipeline Completed Successfully!')

if __name__ == '__main__':
    main()
